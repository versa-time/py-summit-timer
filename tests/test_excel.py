import datetime

from openpyxl import load_workbook

from summit_timer.excel import ExcelRecordWriter
from summit_timer.protocol import DataAck


def make_record(record_number: int) -> DataAck:
    return DataAck(
        device_id=1,
        record_number=record_number,
        event_number=20,
        heat_number=5,
        channel=1,
        record_type="b",
        user_string="45",
        time=datetime.datetime.strptime("01:02:03.4", "%H:%M:%S.%f").time(),
    )


def test_excel_writer_creates_headers_and_appends_records(tmp_path):
    file_path = tmp_path / "records.xlsx"
    writer = ExcelRecordWriter(file_path, "Timing")

    written = writer.append_records([make_record(1), make_record(2)])

    assert [record.record_number for record in written] == [1, 2]

    workbook = load_workbook(file_path)
    sheet = workbook["Timing"]
    assert [cell.value for cell in sheet[1]] == ExcelRecordWriter.headers
    assert sheet.cell(row=2, column=1).value == 1
    assert sheet.cell(row=2, column=2).value == 1
    assert sheet.cell(row=3, column=2).value == 2


def test_excel_writer_does_not_append_duplicate_records(tmp_path):
    file_path = tmp_path / "records.xlsx"
    writer = ExcelRecordWriter(file_path, "Timing")
    writer.append_records([make_record(1)])

    written = ExcelRecordWriter(file_path, "Timing").append_records(
        [make_record(1), make_record(2)]
    )

    workbook = load_workbook(file_path)
    sheet = workbook["Timing"]
    assert [record.record_number for record in written] == [2]
    assert sheet.max_row == 3
