from pathlib import Path
from types import ModuleType
from typing import ClassVar

from openpyxl import Workbook, load_workbook

from summit_timer.protocol import DataAck

TIME_COLUMN_INDEX = 8
TEXT_NUMBER_FORMAT = "@"


class LiveExcelUnavailableError(RuntimeError):
    pass


def format_timer_time(record: DataAck) -> str:
    milliseconds = record.time.microsecond // 1000
    return f"{record.time:%H:%M:%S}.{milliseconds:03d}"


class ExcelRecordWriter:
    headers: ClassVar[list[str]] = [
        "device_id",
        "record_number",
        "event_number",
        "heat_number",
        "channel",
        "record_type",
        "user_string",
        "time",
    ]

    def __init__(self, file_path: str | Path, sheet_name: str):
        self.file_path = Path(file_path)
        self.sheet_name = sheet_name
        self._written_keys: set[tuple[int, int]] = set()

    def append_records(self, records: list[DataAck]) -> list[DataAck]:
        if not records:
            return []

        workbook = self._load_workbook()
        sheet = self._get_sheet(workbook)
        self._ensure_headers(sheet)
        self._load_written_keys(sheet)

        written = []
        for record in records:
            key = (record.device_id, record.record_number)
            if key in self._written_keys:
                continue

            self._append_record(sheet, record)
            self._written_keys.add(key)
            written.append(record)

        if written:
            workbook.save(self.file_path)

        return written

    def _load_workbook(self):
        if self.file_path.exists():
            return load_workbook(self.file_path)
        return Workbook()

    def _get_sheet(self, workbook):
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        sheet = workbook.active
        if self._sheet_has_no_values(sheet):
            sheet.title = self.sheet_name
            return sheet

        return workbook.create_sheet(self.sheet_name)

    def _ensure_headers(self, sheet):
        if self._sheet_has_no_values(sheet):
            self._write_headers(sheet)
            return

        existing_headers = [cell.value for cell in sheet[1][: len(self.headers)]]
        if existing_headers != self.headers:
            raise ValueError(
                f"Sheet '{self.sheet_name}' already exists but does not look like a Summit output sheet."
            )

    def _sheet_has_no_values(self, sheet) -> bool:
        return all(cell.value is None for row in sheet.iter_rows() for cell in row)

    def _write_headers(self, sheet):
        for index, header in enumerate(self.headers, start=1):
            sheet.cell(row=1, column=index, value=header)

    def _append_record(self, sheet, record: DataAck):
        row = max(2, self._last_value_row(sheet) + 1)
        for column, value in enumerate(self._record_values(record), start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            if column == TIME_COLUMN_INDEX:
                cell.number_format = TEXT_NUMBER_FORMAT

    def _record_values(self, record: DataAck) -> list:
        return [
            record.device_id,
            record.record_number,
            record.event_number,
            record.heat_number,
            record.channel,
            record.record_type,
            record.user_string,
            format_timer_time(record),
        ]

    def _last_value_row(self, sheet) -> int:
        for row in range(sheet.max_row, 0, -1):
            if any(cell.value is not None for cell in sheet[row]):
                return row
        return 0

    def _load_written_keys(self, sheet):
        self._written_keys.clear()
        header_map = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        device_column = header_map.get("device_id")
        record_column = header_map.get("record_number")
        if device_column is None or record_column is None:
            return

        for row in range(2, self._last_value_row(sheet) + 1):
            device_id = sheet.cell(row=row, column=device_column).value
            record_number = sheet.cell(row=row, column=record_column).value
            if device_id is None or record_number is None:
                continue
            self._written_keys.add((int(device_id), int(record_number)))


class LiveExcelRecordWriter:
    headers: ClassVar[list[str]] = ExcelRecordWriter.headers

    def __init__(
        self,
        file_path: str | Path,
        sheet_name: str,
        xlwings_module: ModuleType | None = None,
    ):
        self.file_path = Path(file_path)
        self.sheet_name = sheet_name
        self._written_keys: set[tuple[int, int]] = set()
        self.xw = xlwings_module or self._import_xlwings()
        self.book = self._open_or_create_book()

    def append_records(self, records: list[DataAck]) -> list[DataAck]:
        if not records:
            return []

        sheet = self._get_sheet()
        self._ensure_headers(sheet)
        self._load_written_keys(sheet)

        written = []
        rows = []
        for record in records:
            key = (record.device_id, record.record_number)
            if key in self._written_keys:
                continue

            rows.append(self._record_values(record))
            self._written_keys.add(key)
            written.append(record)

        if written:
            start_row = max(2, self._last_value_row(sheet) + 1)
            self._format_time_range_as_text(sheet, start_row, len(rows))
            sheet.range((start_row, 1)).value = rows
            self.book.save(str(self.file_path))

        return written

    def _import_xlwings(self):
        try:
            import xlwings as xw
        except ImportError as exc:
            raise LiveExcelUnavailableError(
                "Live Excel updates require xlwings. Install the GUI extras with "
                "`pdm install -G gui`, then restart the app."
            ) from exc
        return xw

    def _open_or_create_book(self):
        try:
            if self.file_path.exists():
                return self.xw.Book(str(self.file_path))

            book = self.xw.Book()
            book.save(str(self.file_path))
            return book
        except Exception as exc:
            raise LiveExcelUnavailableError(
                f"Could not connect to Excel: {exc}"
            ) from exc

    def _get_sheet(self):
        for sheet in self.book.sheets:
            if sheet.name == self.sheet_name:
                return sheet

        if len(self.book.sheets) == 1 and self._sheet_has_no_values(
            self.book.sheets[0]
        ):
            sheet = self.book.sheets[0]
            sheet.name = self.sheet_name
            return sheet

        return self.book.sheets.add(self.sheet_name, after=self.book.sheets[-1])

    def _ensure_headers(self, sheet):
        if self._sheet_has_no_values(sheet):
            sheet.range((1, 1)).value = self.headers
            return

        existing_headers = self._row_values(sheet, 1, len(self.headers))
        if existing_headers != self.headers:
            raise ValueError(
                f"Sheet '{self.sheet_name}' already exists but does not look like a Summit output sheet."
            )

    def _sheet_has_no_values(self, sheet) -> bool:
        return self._last_value_row(sheet) == 0

    def _last_value_row(self, sheet) -> int:
        used_range = sheet.used_range
        values = used_range.value
        if values is None:
            return 0

        rows = self._as_rows(values)
        for offset, row in enumerate(reversed(rows)):
            if any(value is not None for value in row):
                return used_range.row + len(rows) - offset - 1
        return 0

    def _load_written_keys(self, sheet):
        self._written_keys.clear()
        last_row = self._last_value_row(sheet)
        if last_row < 2:
            return

        for row in range(2, last_row + 1):
            device_id = sheet.range((row, 1)).value
            record_number = sheet.range((row, 2)).value
            if device_id is None or record_number is None:
                continue
            self._written_keys.add((int(device_id), int(record_number)))

    def _record_values(self, record: DataAck) -> list:
        return [
            record.device_id,
            record.record_number,
            record.event_number,
            record.heat_number,
            record.channel,
            record.record_type,
            record.user_string,
            format_timer_time(record),
        ]

    def _format_time_range_as_text(self, sheet, start_row: int, row_count: int):
        end_row = start_row + row_count - 1
        sheet.range(
            (start_row, TIME_COLUMN_INDEX), (end_row, TIME_COLUMN_INDEX)
        ).number_format = TEXT_NUMBER_FORMAT

    def _row_values(self, sheet, row: int, columns: int) -> list:
        values = sheet.range((row, 1), (row, columns)).value
        if columns == 1:
            return [values]
        return values or []

    def _as_rows(self, values) -> list[list]:
        if not isinstance(values, list):
            return [[values]]
        if not values:
            return []
        if not isinstance(values[0], list):
            return [values]
        return values


def create_excel_writer(
    file_path: str | Path,
    sheet_name: str,
    live: bool = True,
) -> LiveExcelRecordWriter | ExcelRecordWriter:
    if live:
        return LiveExcelRecordWriter(file_path, sheet_name)
    return ExcelRecordWriter(file_path, sheet_name)
